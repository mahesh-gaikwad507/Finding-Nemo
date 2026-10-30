"""Admin routes for competition control, live teams monitoring, and Excel export."""

import io
from datetime import datetime, timezone
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy import desc, asc, and_

from database import get_db
from models import User, QuizAttempt, CompetitionSetting
from auth import get_current_user

router = APIRouter(prefix="/api/admin", tags=["Admin"])


def require_admin(current_user: User = Depends(get_current_user)) -> User:
    """Ensure that the current user has admin privileges."""
    is_admin = current_user.is_admin or (current_user.username.lower() in ("mahesh123", "admin"))
    if not is_admin:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Admin access required. Only mahesh123 or administrators can perform this action."
        )
    return current_user


def format_duration(seconds: Optional[int]) -> str:
    """Format seconds into readable format like 2m 15s."""
    if not seconds or seconds <= 0:
        return "—"
    m = seconds // 60
    s = seconds % 60
    h = m // 60
    if h > 0:
        return f"{h}h {m % 60}m {s}s"
    return f"{m}m {s}s"


@router.get("/status")
def get_competition_status(db: Session = Depends(get_db)):
    """Get current competition status (accessible by participants to know when game starts)."""
    setting = db.query(CompetitionSetting).filter(CompetitionSetting.id == 1).first()
    if not setting:
        setting = CompetitionSetting(id=1, is_game_started=False, started_by="mahesh123")
        db.add(setting)
        db.commit()
        db.refresh(setting)

    total_teams = db.query(User).filter(User.is_admin == False).count()
    completed_teams = db.query(QuizAttempt).filter(QuizAttempt.is_completed == True).count()

    return {
        "is_game_started": setting.is_game_started,
        "started_at": setting.started_at,
        "started_by": setting.started_by,
        "total_teams": total_teams,
        "completed_teams": completed_teams
    }


@router.post("/start-game")
def start_competition(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Admin clicks start button: begins the game for all registered teams."""
    setting = db.query(CompetitionSetting).filter(CompetitionSetting.id == 1).first()
    if not setting:
        setting = CompetitionSetting(id=1)
        db.add(setting)

    setting.is_game_started = True
    setting.started_at = datetime.now(timezone.utc)
    setting.started_by = admin.username
    db.commit()
    db.refresh(setting)

    return {
        "message": "Competition started successfully! All teams can now take their quiz.",
        "is_game_started": True,
        "started_at": setting.started_at
    }


@router.post("/reset-game")
def reset_competition(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Admin resets the competition back to waiting mode."""
    setting = db.query(CompetitionSetting).filter(CompetitionSetting.id == 1).first()
    if setting:
        setting.is_game_started = False
        setting.started_at = None
        db.commit()

    return {
        "message": "Competition paused / reset to waiting mode.",
        "is_game_started": False
    }


@router.post("/clear-teams")
def clear_all_teams_history(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Admin clears all history of previously registered teams and resets attempts."""
    try:
        from models import RoomParticipant, RoomTeam, Room, UserAnswer, QuizAttempt
        
        # 1. Remove room participants, room teams, rooms
        db.query(RoomParticipant).delete(synchronize_session=False)
        db.query(RoomTeam).delete(synchronize_session=False)
        db.query(Room).delete(synchronize_session=False)

        # 2. Find all non-admin users
        non_admin_users = db.query(User).filter(
            (User.is_admin == False) & (User.username != "mahesh123")
        ).all()
        non_admin_ids = [u.id for u in non_admin_users]

        if non_admin_ids:
            # Delete user answers
            attempt_ids = [
                a.id for a in db.query(QuizAttempt.id).filter(QuizAttempt.user_id.in_(non_admin_ids)).all()
            ]
            if attempt_ids:
                db.query(UserAnswer).filter(UserAnswer.attempt_id.in_(attempt_ids)).delete(synchronize_session=False)
            
            # Delete quiz attempts
            db.query(QuizAttempt).filter(QuizAttempt.user_id.in_(non_admin_ids)).delete(synchronize_session=False)

            # Delete non-admin users
            db.query(User).filter(User.id.in_(non_admin_ids)).delete(synchronize_session=False)

        # 3. Reset competition setting back to waiting mode
        setting = db.query(CompetitionSetting).filter(CompetitionSetting.id == 1).first()
        if setting:
            setting.is_game_started = False
            setting.started_at = None

        db.commit()

        return {
            "message": f"Successfully cleared all history of previously registered players ({len(non_admin_ids)} teams removed). Ready for new competition!",
            "cleared_count": len(non_admin_ids),
            "is_game_started": False
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to clear history: {str(e)}")


@router.get("/teams")
def get_registered_teams(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Get all registered teams with real-time quiz results for admin live monitor."""
    users = db.query(User).filter(User.is_admin == False).order_by(User.id.asc()).all()

    teams_data = []
    for u in users:
        # Find latest/best attempt
        attempt = db.query(QuizAttempt).filter(
            QuizAttempt.user_id == u.id
        ).order_by(
            QuizAttempt.is_completed.desc(),
            QuizAttempt.score.desc(),
            QuizAttempt.id.desc()
        ).first()

        status_label = "Waiting"
        domain = u.locked_domain or "Not Selected"
        score = 0
        correct = 0
        wrong = 0
        total_time = 0
        speed_str = "—"
        completed_at = None

        if attempt:
            domain = attempt.category.upper()
            score = attempt.score
            correct = attempt.correct_answers
            wrong = attempt.wrong_answers
            total_time = attempt.total_time_seconds or 0
            answered = correct + wrong
            if answered > 0 and total_time > 0:
                speed_str = f"{(total_time / answered):.1f}s / Q"
            if attempt.is_completed:
                status_label = "Completed"
                completed_at = attempt.completed_at
            else:
                status_label = f"In Progress (Q{attempt.current_question})"

        teams_data.append({
            "user_id": u.id,
            "username": u.username,
            "team_name": u.team_name or u.username,
            "leader_name": u.leader_name or "—",
            "phone_number": u.phone_number or "—",
            "domain": domain,
            "score": score,
            "correct_answers": correct,
            "wrong_answers": wrong,
            "total_time_seconds": total_time,
            "time_formatted": format_duration(total_time),
            "answering_speed": speed_str,
            "status": status_label,
            "completed_at": completed_at.strftime("%Y-%m-%d %H:%M:%S") if completed_at else "—"
        })

    # Sort by Score (DESC) -> Correct Answers (DESC) -> Total Time (ASC)
    teams_data.sort(
        key=lambda t: (
            -t["score"],
            -t["correct_answers"],
            t["total_time_seconds"] if t["total_time_seconds"] > 0 else 999999
        )
    )

    # Assign ranks
    for rank, t in enumerate(teams_data, 1):
        t["rank"] = rank

    return teams_data


@router.get("/export-excel")
def export_teams_excel(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """Generate and download real Microsoft Excel (.xlsx) file with all team records."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Competition Scorecard"

    # Title Row
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "Techflix — Finding Nemo: Official Scorecard"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Generated Timestamp
    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | Organizer: mahesh123"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "Rank", "Team Name", "Leader Name", "Phone No.",
        "Field / Domain", "Score", "Correct Qs", "Mistakes",
        "Time Taken", "Answering Speed", "Status"
    ]
    ws.append([])  # blank row 3
    ws.append(headers)  # row 4
    ws.row_dimensions[4].height = 26

    header_fill = PatternFill(start_color="0369A1", end_color="0369A1", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Fetch teams and rank them
    users = db.query(User).filter(User.is_admin == False).all()
    rows_data = []

    for u in users:
        attempt = db.query(QuizAttempt).filter(
            QuizAttempt.user_id == u.id
        ).order_by(
            QuizAttempt.is_completed.desc(),
            QuizAttempt.score.desc(),
            QuizAttempt.id.desc()
        ).first()

        status_label = "Waiting"
        domain = u.locked_domain or "Not Selected"
        score = 0
        correct = 0
        wrong = 0
        total_time = 0
        speed_str = "—"

        if attempt:
            domain = attempt.category.upper()
            score = attempt.score
            correct = attempt.correct_answers
            wrong = attempt.wrong_answers
            total_time = attempt.total_time_seconds or 0
            answered = correct + wrong
            if answered > 0 and total_time > 0:
                speed_str = f"{(total_time / answered):.1f}s / Q"
            status_label = "Completed" if attempt.is_completed else f"In Progress (Q{attempt.current_question})"

        rows_data.append({
            "team_name": u.team_name or u.username,
            "leader_name": u.leader_name or "—",
            "phone_number": u.phone_number or "—",
            "domain": domain,
            "score": score,
            "correct": correct,
            "wrong": wrong,
            "total_time": total_time,
            "time_str": format_duration(total_time),
            "speed_str": speed_str,
            "status": status_label
        })

    # Sort
    rows_data.sort(
        key=lambda r: (
            -r["score"],
            -r["correct"],
            r["total_time"] if r["total_time"] > 0 else 999999
        )
    )

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for rank, r in enumerate(rows_data, 1):
        row_vals = [
            rank,
            r["team_name"],
            r["leader_name"],
            r["phone_number"],
            r["domain"],
            r["score"],
            f"{r['correct']} / 30",
            r["wrong"],
            r["time_str"],
            r["speed_str"],
            r["status"]
        ]
        ws.append(row_vals)
        curr_row = ws.max_row
        ws.row_dimensions[curr_row].height = 22

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.font = Font(name="Calibri", size=11)
            c.border = thin_border
            if rank % 2 == 0:
                c.fill = alt_fill
            if col_idx in (1, 6, 7, 8, 9, 10, 11):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    # Save to buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Techflix_Competition_Records_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )
